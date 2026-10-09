from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
root=Path(__file__).resolve().parents[1]/'templates'
for kind,cols in [('musteri',['ad','soyad','telefon','şehir','mağaza','müşteri grubu','son alışveriş tarihi']),('personel',['ad','soyad','telefon','mağaza','departman','görev'])]:
    wb=Workbook();ws=wb.active;ws.title='Aktarım';ws.append(cols)
    for cell in ws[1]:cell.font=Font(color='FFFFFF',bold=True);cell.fill=PatternFill('solid',fgColor='267652');cell.alignment=Alignment(vertical='center')
    ws.row_dimensions[1].height=28;ws.freeze_panes='A2'
    for col in ws.columns:ws.column_dimensions[col[0].column_letter].width=24
    for row in ws.iter_rows(min_row=2,max_row=1001,min_col=3,max_col=3):row[0].number_format='@'
    info=wb.create_sheet('Talimat');info.append(['Telefonları metin olarak girin. XLSX veya UTF-8 CSV yükleyin.']);info.append(['İzin bilgileri bu dosyadan doğrulanmaz. Panelde kanıtı inceleyerek ayrı doğrulayın.']);info.column_dimensions['A'].width=110
    wb.save(root/f'{kind}_aktarim_sablonu.xlsx')
