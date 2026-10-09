import csv, io, json, secrets, zipfile
from pathlib import Path
from datetime import datetime, timedelta
from fastapi import Depends, HTTPException, UploadFile, File
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field
from sqlalchemy import select, func, or_
from app.core import *
from app.main import app, actor, admin, db_session, serial, require

class ContactInput(BaseModel):
    first_name:str=Field(min_length=1,max_length=100)
    last_name:str=Field(default='',max_length=100)
    phone:str
    city:str=''
    store:str=''
    group:str=''
    department:str=''
    job:str=''
    last_purchase:str=''
    active:bool=True
    preference:str='whatsapp'
@app.get('/api/contacts/{kind}')
def contacts(kind:str,q:str='',city:str='',store:str='',group:str='',consent:str='',page:int=1,user=Depends(actor),db=Depends(db_session)):
    stmt=select(Contact).where(Contact.kind==kind)
    if q:stmt=stmt.where(or_(Contact.first_name.ilike(f'%{q}%'),Contact.last_name.ilike(f'%{q}%'),Contact.phone.ilike(f'%{q}%')))
    for field,value in [(Contact.city,city),(Contact.store,store),(Contact.group,group),(Contact.consent,consent)]:
        if value:stmt=stmt.where(field==value)
    total=db.scalar(select(func.count()).select_from(stmt.subquery()))
    rows=db.scalars(stmt.order_by(Contact.id.desc()).offset((max(1,page)-1)*50).limit(50)).all()
    audit(db,user.username,'contact_access',f'{kind}; sayfa {page}');db.commit()
    return {'total':total,'items':[{**serial(c),'phone':c.phone if user.role=='admin' else mask(c.phone),'eligible':eligible(c)} for c in rows]}
@app.post('/api/contacts/{kind}')
def create_contact(kind:str,data:ContactInput,user=Depends(actor),db=Depends(db_session)):
    if kind not in ('customer','staff'):raise HTTPException(422,'Geçersiz kayıt türü')
    values=data.model_dump()
    try:values['phone']=phone(values['phone'])
    except ValueError as e:raise HTTPException(422,str(e))
    if db.scalar(select(Contact.id).where(Contact.kind==kind,Contact.phone==values['phone'])):raise HTTPException(409,'Telefon zaten kayıtlı')
    c=Contact(kind=kind,**values);db.add(c);audit(db,user.username,'contact_create',kind);db.commit();return {'id':c.id}
@app.put('/api/contact/{id}')
def edit_contact(id:int,data:ContactInput,user=Depends(admin),db=Depends(db_session)):
    c=require(db,Contact,id);values=data.model_dump()
    try:values['phone']=phone(values['phone'])
    except ValueError as e:raise HTTPException(422,str(e))
    if values['phone']!=c.phone:
        if db.scalar(select(Contact.id).where(Contact.kind==c.kind,Contact.phone==values['phone'])):raise HTTPException(409,'Telefon zaten kayıtlı')
        c.consent='unverified';c.iys_status='not_checked';c.evidence='';c.opted_out=True
    for k,v in values.items():setattr(c,k,v)
    audit(db,user.username,'contact_update',str(id));db.commit();return {'ok':True}
@app.delete('/api/contact/{id}')
def delete_contact(id:int,user=Depends(admin),db=Depends(db_session)):
    c=require(db,Contact,id)
    for m in db.scalars(select(Message).where(Message.contact_id==id)):
        m.contact_id=None;m.phone='deleted';m.body='Kişisel veri silindi'
        if m.status in ('queued','retry'):m.status='cancelled'
    db.delete(c);audit(db,user.username,'personal_data_delete',str(id));db.commit();return {'ok':True}
class ConsentInput(BaseModel):
    verified:bool
    scope:str=''
    source:str=''
    evidence:str=''
    consent_date:str=''
    legal_basis:str=''
    iys_manual_verified:bool=False
    opted_out:bool=False
@app.put('/api/contact/{id}/consent')
def consent(id:int,data:ConsentInput,user=Depends(admin),db=Depends(db_session)):
    c=require(db,Contact,id)
    if data.verified:
        if not all([data.source.strip(),data.evidence.strip(),data.consent_date.strip()]):raise HTTPException(422,'İzin kaynağı, kanıtı ve tarihi zorunlu')
        try:datetime.fromisoformat(data.consent_date)
        except ValueError:raise HTTPException(422,'İzin tarihi ISO biçiminde olmalı')
        if data.scope!=('marketing' if c.kind=='customer' else 'staff'):raise HTTPException(422,'İzin kapsamı uygun değil')
        if c.kind=='staff' and not data.legal_basis.strip():raise HTTPException(422,'Personel iletişim dayanağı zorunlu')
    c.consent='verified' if data.verified and not data.opted_out else 'revoked' if data.opted_out else 'unverified'
    for field in ('scope','source','evidence','consent_date','legal_basis','opted_out'):setattr(c,field,getattr(data,field))
    c.iys_status='manual_verified' if data.iys_manual_verified else 'not_checked'
    audit(db,user.username,'consent_change',json.dumps({'id':id,**data.model_dump()},ensure_ascii=False));db.commit()
    return {'eligible':eligible(c),'iys':'Manuel belge kontrolü; API sorgusu yapılmadı'}
ALIASES={'ad':'first_name','soyad':'last_name','telefon':'phone','şehir':'city','sehir':'city','mağaza':'store','magaza':'store','müşteri grubu':'group','musteri_grubu':'group','departman':'department','görev':'job','gorev':'job','son alışveriş tarihi':'last_purchase'}
@app.post('/api/import/{kind}/preview')
async def import_preview(kind:str,file:UploadFile=File(...),user=Depends(actor),db=Depends(db_session)):
    if kind not in ('customer','staff'):raise HTTPException(422,'Geçersiz tür')
    raw=await file.read(5*1024*1024+1)
    if len(raw)>5*1024*1024:raise HTTPException(413,'Dosya en fazla 5 MB olabilir')
    try:
        ext=Path(file.filename or '').suffix.lower()
        if ext=='.xlsx':
            with zipfile.ZipFile(io.BytesIO(raw)) as z:
                if sum(i.file_size for i in z.infolist())>40*1024*1024:raise ValueError('Açılmış dosya boyutu çok büyük')
            from openpyxl import load_workbook
            workbook=load_workbook(io.BytesIO(raw),read_only=True,data_only=True)
            rows=[]
            for row in workbook.active.iter_rows(values_only=True):
                rows.append(row)
                if len(rows)>10001:raise ValueError('En fazla 10.000 satır yükleyin')
            workbook.close()
        elif ext=='.csv':
            content=raw.decode('utf-8-sig')
            try:dialect=csv.Sniffer().sniff(content[:4096],delimiters=',;\t')
            except csv.Error:dialect=csv.excel
            rows=list(csv.reader(io.StringIO(content),dialect))
        else:raise ValueError('Yalnızca XLSX ve UTF-8 CSV desteklenir')
        if len(rows)>10001:raise ValueError('En fazla 10.000 satır yükleyin')
        if not rows:raise ValueError('Dosya boş')
        headers=[ALIASES.get(str(h).strip().lower(),str(h).strip().lower()) for h in rows[0]]
        if not {'first_name','phone'}.issubset(headers):raise ValueError('ad ve telefon sütunları zorunlu')
        existing=set(db.scalars(select(Contact.phone).where(Contact.kind==kind)));seen=set();valid=[];errors=[]
        for index,row in enumerate(rows[1:],2):
            if not any(v is not None and str(v).strip() for v in row):continue
            try:
                values={k:str(v).strip() if v is not None else '' for k,v in zip(headers,row) if k in ContactInput.model_fields and k not in ('active','preference')}
                values['phone']=phone(values.get('phone',''));record=ContactInput(**values).model_dump()
                if record['phone'] in seen or record['phone'] in existing:raise ValueError('Mükerrer telefon')
                seen.add(record['phone']);valid.append(record)
            except Exception as e:errors.append({'row':index,'error':str(e)[:250]})
        token=secrets.token_urlsafe(32)
        db.add(SystemValue(key='import:'+token,value=json.dumps({'user':user.id,'kind':kind,'rows':valid,'expires':(now()+timedelta(minutes=30)).isoformat()})));db.commit()
        return {'token':token,'valid_count':len(valid),'errors':errors,'preview':valid[:20],'notice':'Excel izin bilgileri doğrulanmış izin kabul edilmez. Kayıtlar inceleme bekliyor olarak aktarılır.'}
    except HTTPException:raise
    except Exception as e:raise HTTPException(422,str(e))
class ImportCommit(BaseModel):token:str
@app.post('/api/import/{kind}/commit')
def import_commit(kind:str,data:ImportCommit,user=Depends(actor),db=Depends(db_session)):
    batch=require(db,SystemValue,'import:'+data.token);payload=json.loads(batch.value)
    if payload['user']!=user.id or payload['kind']!=kind or datetime.fromisoformat(payload['expires'])<now():raise HTTPException(403,'Önizleme süresi dolmuş veya yetkisiz')
    count=0;duplicates=0
    for row in payload['rows']:
        if db.scalar(select(Contact.id).where(Contact.kind==kind,Contact.phone==row['phone'])):duplicates+=1;continue
        db.add(Contact(kind=kind,**row));count+=1
    db.delete(batch);audit(db,user.username,'import',f'{kind}: {count}');db.commit();return {'imported':count,'duplicates':duplicates}
def safe_cell(v):return "'"+v if isinstance(v,str) and v.startswith(('=','+','-','@')) else v
@app.get('/api/export/{kind}')
def export_contacts(kind:str,user=Depends(admin),db=Depends(db_session)):
    from openpyxl import Workbook
    wb=Workbook();ws=wb.active;ws.title='Kayıtlar'
    cols=['first_name','last_name','phone','city','store','group','department','job','consent','iys_status'];ws.append(cols)
    for c in db.scalars(select(Contact).where(Contact.kind==kind)):ws.append([safe_cell(getattr(c,k)) for k in cols])
    buf=io.BytesIO();wb.save(buf);buf.seek(0);audit(db,user.username,'export',kind);db.commit()
    return StreamingResponse(buf,media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':f'attachment; filename="{kind}.xlsx"'})
@app.get('/api/stores')
def stores(user=Depends(actor),db=Depends(db_session)):return [serial(s) for s in db.scalars(select(Store).order_by(Store.name))]
class StoreInput(BaseModel):
    name:str=Field(min_length=1,max_length=160)
    city:str=''
@app.post('/api/stores')
def store_add(data:StoreInput,user=Depends(admin),db=Depends(db_session)):
    if db.scalar(select(Store).where(Store.name==data.name)):raise HTTPException(409,'Mağaza mevcut')
    s=Store(**data.model_dump());db.add(s);audit(db,user.username,'store_create',data.name);db.commit();return serial(s)
