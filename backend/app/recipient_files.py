"""Bounded spreadsheet readers shared by audience and external permission imports."""
import csv, io, math, re, unicodedata, zipfile
from pathlib import Path
from fastapi import HTTPException
from app.core import phone

ALIASES={'phone','telephone','mobile','gsm','cep','telefon','telefonnumarasi','telefonno','ceptel','mobilephone','ceptelefonu'}

def cell_phone(value):
    if isinstance(value,float):
        if not math.isfinite(value) or not value.is_integer():return None
        value=int(value)
    try:
        result=phone(str(value or '').strip())
        return result if re.fullmatch(r'\+905\d{9}',result) else None
    except ValueError:return None

def header(value):
    value=str(value or '').casefold().replace('ı','i')
    return ''.join(c for c in unicodedata.normalize('NFKD',value) if c.isalnum() and not unicodedata.combining(c))

def read_rows(raw,filename):
    if len(raw)>5*1024*1024:raise HTTPException(413,'Dosya en fazla 5 MB olabilir.')
    try:
        ext=Path(filename or '').suffix.lower()
        if ext=='.xlsx':
            from openpyxl import load_workbook
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                if sum(i.file_size for i in archive.infolist())>40*1024*1024:raise ValueError('Açılmış Excel dosyası çok büyük.')
            wb=load_workbook(io.BytesIO(raw),read_only=True,data_only=True)
            try:
                if wb.active.max_column>100:raise ValueError('En fazla 100 sütun yükleyin.')
                rows=[]
                for row in wb.active.iter_rows(values_only=True):
                    rows.append(list(row))
                    if len(rows)>50001:raise ValueError('Dosya satır sınırını aşıyor.')
            finally:wb.close()
        elif ext=='.xls':
            import xlrd
            wb=xlrd.open_workbook(file_contents=raw,on_demand=True)
            try:
                sheet=wb.sheet_by_index(0)
                if sheet.nrows>50001 or sheet.ncols>100:raise ValueError('Dosya satır veya sütun sınırını aşıyor.')
                rows=[sheet.row_values(i) for i in range(sheet.nrows)]
            finally:wb.release_resources()
        elif ext=='.csv':
            try:text=raw.decode('utf-8-sig')
            except UnicodeDecodeError:text=raw.decode('cp1254')
            try:dialect=csv.Sniffer().sniff(text[:4096],delimiters=',;\t')
            except csv.Error:dialect=csv.excel
            rows=[]
            for row in csv.reader(io.StringIO(text),dialect):
                if len(row)>100:raise ValueError('En fazla 100 sütun yükleyin.')
                rows.append(row)
                if len(rows)>50001:raise ValueError('Dosya satır sınırını aşıyor.')
        else:raise ValueError('XLSX, XLS veya CSV dosyası seçin.')
        if not rows:raise ValueError('Dosya boş.')
        return rows
    except HTTPException:raise
    except Exception as e:raise HTTPException(422,str(e) if isinstance(e,ValueError) else 'Dosya okunamadı. Excel/CSV biçimini kontrol edin.')

def parse_recipients(rows,selected=None):
    first=next((i for i,row in enumerate(rows) if any(str(v or '').strip() for v in row)),None)
    if first is None:raise HTTPException(422,'Dosyada dolu satır bulunamadı.')
    # Headerless phone lists retain their first recipient.
    has_header=any(header(v) in ALIASES or header(v).startswith('telefon') for v in rows[first]) or not any(cell_phone(v) for v in rows[first])
    width=max(map(len,rows))
    headers=[str(v or f'Sütun {i+1}') for i,v in enumerate(rows[first])] if has_header else [f'Sütun {i+1}' for i in range(width)]
    headers+= [f'Sütun {i+1}' for i in range(len(headers),width)]
    data=list(enumerate(rows[first+1:] if has_header else rows[first:],first+2 if has_header else first+1))
    named=[i for i,h in enumerate(headers) if header(h) in ALIASES or header(h).startswith('telefon')]
    sample=[r for _,r in data if any(str(v or '').strip() for v in r)][:100]
    scores=[sum(cell_phone(row[i]) is not None for row in sample if i<len(row)) for i in range(width)]
    best=max(scores,default=0)
    candidates=[i for i,s in enumerate(scores) if s==best and best>0]
    column=selected if selected is not None else named[0] if len(named)==1 else candidates[0] if len(candidates)==1 else None
    if column is not None and not 0<=column<width:raise HTTPException(422,'Telefon sütunu geçersiz.')
    numbers=[];seen=set();invalid=[];duplicates=skipped=0
    for index,row in data:
        if not any(v is not None and str(v).strip() for v in row):skipped+=1;continue
        if len(numbers)+len(invalid)+duplicates>=10000:raise HTTPException(422,'En fazla 10.000 dolu alıcı satırı yükleyin.')
        if column is None:continue
        value=row[column] if column<len(row) else None
        normalized=cell_phone(value)
        if not normalized:invalid.append({'row':index,'value':str(value or '')[:80],'reason':'Geçersiz Türk cep telefonu numarası'});continue
        if normalized in seen:duplicates+=1;continue
        seen.add(normalized);numbers.append(normalized)
    return {'numbers':numbers,'invalid':invalid,'duplicate_count':duplicates,'skipped_count':skipped,'phone_column':column,'columns':[{'index':i,'name':h} for i,h in enumerate(headers)],'needs_column':column is None}
