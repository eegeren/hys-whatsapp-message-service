"""Permission-gated bulk campaign preparation using existing customer records."""
import csv,io,json,secrets,zipfile
from datetime import timedelta
from pathlib import Path
from fastapi import Depends,File,HTTPException,UploadFile
from pydantic import BaseModel,Field
from sqlalchemy import select,func
from app.core import settings,Contact,Template,Campaign,Message,SystemValue,now,audit,phone,eligible,mask
from app.main import app,actor,admin,db_session,require,connection_ok

ALIASES={'phone','telephone','mobile','gsm','cep','telefon','telefonnumarasi','telefonno','ceptel','mobilephone'}
def cell_phone(value):
    if value is None:return None
    try:return phone(str(value).strip())
    except ValueError:return None

@app.post('/api/bulk/preview')
async def bulk_preview(file:UploadFile=File(...),user=Depends(actor),db=Depends(db_session)):
    raw=await file.read(5*1024*1024+1)
    if len(raw)>5*1024*1024:raise HTTPException(413,'Dosya en fazla 5 MB olabilir.')
    try:
        ext=Path(file.filename or '').suffix.lower()
        if ext=='.xlsx':
            from openpyxl import load_workbook
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                if sum(i.file_size for i in archive.infolist())>40*1024*1024:raise ValueError('Açılmış Excel dosyası sınırı aşıyor.')
            wb=load_workbook(io.BytesIO(raw),read_only=True,data_only=True);rows=list(wb.active.iter_rows(values_only=True));wb.close()
        elif ext=='.csv':
            text=raw.decode('utf-8-sig')
            try:dialect=csv.Sniffer().sniff(text[:4096],delimiters=',;\t')
            except csv.Error:dialect=csv.excel
            rows=list(csv.reader(io.StringIO(text),dialect))
        else:raise ValueError('Yalnızca XLSX ve UTF-8 CSV desteklenir.')
        if not rows or len(rows)>10001:raise ValueError('Dosya boş veya satır sınırını aşıyor.')
        headers=[''.join(ch for ch in str(x or '').casefold() if ch.isalnum()) for x in rows[0]]
        phone_col=next((i for i,h in enumerate(headers) if h in ALIASES or h.startswith('telefon')),None)
        data_rows=rows[1:]
        if phone_col is None:
            scores=[]
            for col in range(max(map(len,data_rows),default=0)):
                values=[r[col] if col<len(r) else None for r in data_rows[:25]]
                scores.append((sum(cell_phone(v) is not None for v in values),col))
            score,phone_col=max(scores,default=(0,-1))
            if score<1:raise ValueError('Telefon numarası kolonu bulunamadı.')
        numbers=[];seen=set();duplicates=invalid=0
        skipped=0
        for row in data_rows:
            # Empty worksheet/CSV rows are layout noise, not invalid recipients.
            if not row or not any(str(value or '').strip() for value in row):
                skipped+=1
                continue
            value=row[phone_col] if phone_col<len(row) else None
            normalized=cell_phone(value)
            if not normalized:invalid+=1;continue
            if normalized in seen:duplicates+=1;continue
            seen.add(normalized);numbers.append(normalized)
    except HTTPException:raise
    except Exception as e:raise HTTPException(422,str(e) if isinstance(e,ValueError) else 'Dosya okunamadı; CSV/XLSX biçimini kontrol edin.')
    contacts=db.scalars(select(Contact).where(Contact.kind=='customer',Contact.phone.in_(numbers))).all() if numbers else []
    by_phone={c.phone:c for c in contacts};allowed=[by_phone[n].id for n in numbers if n in by_phone and eligible(by_phone[n])]
    token=secrets.token_urlsafe(32)
    db.add(SystemValue(key='bulk_preview:'+token,value=json.dumps({'user':user.id,'ids':allowed,'expires':(now()+timedelta(minutes=30)).isoformat()})))
    audit(db,user.username,'bulk_file_preview',json.dumps({'unique':len(numbers),'eligible':len(allowed),'duplicates':duplicates,'invalid':invalid}));db.commit()
    return {'token':token,'unique_count':len(numbers),'duplicate_count':duplicates,'invalid_count':invalid,'skipped_count':skipped,'eligible_count':len(allowed),'excluded_count':len(numbers)-len(allowed),'eligible_preview':[{'phone':mask(c.phone)} for c in contacts if c.id in allowed[:5]],'cost':None,'cost_notice':'Meta fiyatı için etkin tarife kaynağı yok; tahmini maliyet bilinmiyor.','permission_notice':'Yalnızca mevcut HYS kayıtlarında doğrulanmış pazarlama izni bulunan müşteriler seçildi. Dosyadaki izin sütunları tek başına izin sayılmaz.','expires_in_minutes':30}

class BulkCampaignInput(BaseModel):
    token:str=Field(min_length=30,max_length=100)
    template_id:int
    variables:dict=Field(default_factory=dict)

@app.post('/api/bulk/campaign')
def create_bulk_campaign(data:BulkCampaignInput,user=Depends(actor),db=Depends(db_session)):
    saved=require(db,SystemValue,'bulk_preview:'+data.token);payload=json.loads(saved.value)
    if payload.get('user')!=user.id or datetime_from(payload.get('expires'))<now():raise HTTPException(403,'Toplu mesaj önizlemesinin süresi dolmuş; dosyayı yeniden yükleyin.')
    if not payload.get('ids'):raise HTTPException(422,'Doğrulanmış izinli alıcı bulunamadı; kampanya oluşturulmadı.')
    template=require(db,Template,data.template_id)
    if template.status!='APPROVED' or template.category!='MARKETING':raise HTTPException(403,'Reklam gönderimi için Meta onaylı Marketing şablonu zorunludur.')
    for cid in payload['ids']:
        contact=require(db,Contact,cid)
        if contact.kind!='customer' or not eligible(contact):raise HTTPException(403,'Alıcı izinleri değişti; dosyayı yeniden önizleyin.')
    campaign=Campaign(kind='customer',name='Toplu mesaj · '+template.name,description='İzinli XLSX/CSV alıcı listesinden oluşturuldu.',template_id=template.id,filters=json.dumps({'ids':payload['ids']}),variables=json.dumps(data.variables),max_count=len(payload['ids']),budget='')
    db.add(campaign);audit(db,user.username,'bulk_campaign_create',str(len(payload['ids'])));db.commit();db.refresh(campaign)
    return {'campaign':{'id':campaign.id,'name':campaign.name,'status':campaign.status,'count':len(payload['ids'])},'dry_run':settings.dry_run,'live_send_enabled':settings.live_send_enabled}

def datetime_from(value):
    from datetime import datetime
    return datetime.fromisoformat(value)

@app.get('/api/bulk/campaign/{campaign_id}')
def bulk_campaign_status(campaign_id:int,user=Depends(actor),db=Depends(db_session)):
    campaign=require(db,Campaign,campaign_id)
    if campaign.kind!='customer':raise HTTPException(404,'Toplu müşteri kampanyası bulunamadı.')
    rows=db.scalars(select(Message).where(Message.campaign_id==campaign_id)).all()
    return {'campaign':{'id':campaign.id,'name':campaign.name,'status':campaign.status,'first_approval':campaign.first_approval,'second_approval':campaign.second_approval},'total':len(rows),'queued':sum(m.status in ('queued','retry','sending') for m in rows),'dry_run':sum(m.status=='dry_run' for m in rows),'accepted':sum(m.status in ('accepted','sent','delivered','read') for m in rows),'delivered':sum(m.status in ('delivered','read') for m in rows),'read':sum(m.status=='read' for m in rows),'failed':sum(m.status in ('blocked','failed','uncertain') for m in rows),'dry_run_enabled':settings.dry_run,'live_send_enabled':settings.live_send_enabled,'connection_verified':connection_ok(db)}

from datetime import datetime
