import io,json,hmac,hashlib
from datetime import timedelta
import pytest
from sqlalchemy import select
from openpyxl import Workbook,load_workbook
from fastapi import HTTPException
from app.core import *
from app.worker import process_one,finalize,reserve_bulk_slot
from app.recipient_files import read_rows,parse_recipients

def upload(client,body='Telefon\n05321234567\n05331234567\n'):
    r=client.post('/api/bulk/preview',files={'file':('numbers.csv',body.encode(),'text/csv')})
    assert r.status_code==200,r.text
    return r.json()

def permit(client,body='Telefon\n05321234567\n05331234567\n',confirmed=True):
    return client.post('/api/bulk/permissions/import',files={'file':('permission.csv',body.encode(),'text/csv')},data={'confirmed':str(confirmed).lower()})

def template():
    with Session() as db:
        t=Template(name='bulk_marketing',status='APPROVED',meta_id='777',category='MARKETING',language='tr',body='Merhaba {{1}}',components=json.dumps([{'type':'BODY','text':'Merhaba {{1}}'}]))
        db.add(t);db.commit();return t.id

def prepare(client):
    assert permit(client).status_code==200
    p=upload(client);tid=template()
    r=client.post('/api/bulk/campaign',json={'token':p['token'],'template_id':tid,'variables':{'1':'HYS'}})
    assert r.status_code==200,r.text
    return r.json()['campaign']['id'],p,tid

def fake_live(monkeypatch):
    import app.bulk as bulk,app.worker as worker,app.main as main
    monkeypatch.setattr(settings,'dry_run',False);monkeypatch.setattr(settings,'live_send_enabled',True)
    monkeypatch.setattr(bulk,'connection_ok',lambda db:True);monkeypatch.setattr(main,'connection_ok',lambda db:True)
    def lookup(*args,**kwargs):
        with Session() as db:
            t=db.scalar(select(Template))
            return {'data':[{'id':t.meta_id,'name':t.name,'status':t.status,'category':t.category,'language':t.language,'components':json.loads(t.components)}]}
    monkeypatch.setattr(bulk,'graph',lookup)
    monkeypatch.setattr(worker,'reserve_bulk_slot',lambda db:True)
    calls=[]
    def send(method,path,payload=None,**kwargs):
        assert method=='POST' and path==settings.meta_phone_number_id+'/messages'
        calls.append(payload);return {'messages':[{'id':f'wamid.bulk.{len(calls)}'}]}
    monkeypatch.setattr(worker,'graph',send);return calls

def test_permission_export_requires_attestation_and_never_assumes_main_file(client):
    p=upload(client);assert p['eligible_count']==0
    assert permit(client,confirmed=False).status_code==422
    assert permit(client).json()['imported']==2
    p=upload(client);assert p['eligible_count']==2
    with Session() as db:
        c=db.scalar(select(Contact));assert c.first_name=='';assert not c.consent_date
        assert c.iys_status=='external_verified';assert eligible(c)
        assert 'external_marketing_export:' in c.evidence

def test_permission_import_never_overrides_cross_module_optout(client):
    with Session() as db:db.add(Contact(kind='staff',first_name='Ret',phone='+905321234567',opted_out=True));db.commit()
    r=permit(client).json();assert r['blocked']==1;assert r['imported']==1
    p=upload(client);assert p['eligible_count']==1;assert p['excluded_count']==1

def test_excel_phone_only_499_blank_rows_float_and_duplicate(client):
    wb=Workbook();ws=wb.active;ws.append(['Telefon']);ws.append([5321234567.0])
    for _ in range(499):ws.append([None])
    ws.append(['+905321234567']);ws.append(['abc'])
    b=io.BytesIO();wb.save(b)
    r=client.post('/api/bulk/preview',files={'file':('numbers.xlsx',b.getvalue(),'application/octet-stream')})
    assert r.status_code==200,r.text
    p=r.json();assert p['skipped_count']==499;assert p['unique_count']==1;assert p['duplicate_count']==1;assert p['invalid_count']==1

def test_headerless_preserves_first_recipient_and_ambiguous_column_selection(client):
    assert upload(client,'05321234567\n05331234567\n')['unique_count']==2
    body='A,B\n05321234567,05331234567\n'
    p=upload(client,body);assert p['needs_column']
    r=client.post('/api/bulk/preview',files={'file':('numbers.csv',body.encode(),'text/csv')},data={'phone_column':'1'})
    assert r.json()['unique_count']==1;assert r.json()['phone_column']==1

def test_xls_reader_fixture():
    from pathlib import Path
    raw=(Path(__file__).parent/'fixtures'/'recipients.xls').read_bytes()
    p=parse_recipients(read_rows(raw,'recipients.xls'))
    assert p['numbers']==['+905321234567','+905331234567']
    assert p['duplicate_count']==1

def test_confirmed_single_admin_start_and_double_click_never_duplicates(client,monkeypatch):
    cid,p,tid=prepare(client);calls=fake_live(monkeypatch)
    assert client.post(f'/api/bulk/campaign/{cid}/start',json={'confirmed':False}).status_code==422
    assert calls==[]
    r=client.post(f'/api/bulk/campaign/{cid}/start',json={'confirmed':True});assert r.json()['queued']==2
    assert client.post(f'/api/bulk/campaign/{cid}/start',json={'confirmed':True}).status_code==409
    with Session() as db:ids=list(db.scalars(select(Message.id)))
    for mid in ids:process_one(mid);process_one(mid)
    assert len(calls)==2;assert len({c['to'] for c in calls})==2
    p=client.get(f'/api/bulk/campaign/{cid}').json()
    assert p['accepted']==2;assert p['delivered']==0;assert p['sent']==0;assert p['read']==0

def test_sample_then_bulk_never_resends_sample(client,monkeypatch):
    cid,_,_=prepare(client);calls=fake_live(monkeypatch)
    assert client.post(f'/api/bulk/campaign/{cid}/test',json={'confirmed':True,'phone':'05321234567'}).json()['queued']==1
    assert client.post(f'/api/bulk/campaign/{cid}/start',json={'confirmed':True}).status_code==409
    with Session() as db:mid=db.scalar(select(Message.id))
    process_one(mid);finalize()
    r=client.post(f'/api/bulk/campaign/{cid}/start',json={'confirmed':True});assert r.json()['queued']==1
    with Session() as db:ids=list(db.scalars(select(Message.id)))
    for mid in ids:process_one(mid)
    assert len(calls)==2;assert calls[0]['to']!=calls[1]['to']

def test_pause_resume_cancel_preserves_successful_records(client,monkeypatch):
    cid,_,_=prepare(client);calls=fake_live(monkeypatch)
    client.post(f'/api/bulk/campaign/{cid}/start',json={'confirmed':True})
    with Session() as db:ids=list(db.scalars(select(Message.id)))
    process_one(ids[0]);assert len(calls)==1
    assert client.post(f'/api/campaigns/{cid}/pause').status_code==200
    process_one(ids[1]);assert len(calls)==1
    assert client.post(f'/api/campaigns/{cid}/resume').status_code==200
    process_one(ids[0]);assert len(calls)==1
    assert client.post(f'/api/campaigns/{cid}/cancel').status_code==200
    process_one(ids[1]);assert len(calls)==1
    with Session() as db:assert db.get(Message,ids[1]).status=='cancelled'

def test_runtime_optout_and_sender_change_block_without_post(client,monkeypatch):
    cid,_,_=prepare(client);calls=fake_live(monkeypatch)
    client.post(f'/api/bulk/campaign/{cid}/start',json={'confirmed':True})
    with Session() as db:
        rows=db.scalars(select(Message)).all();ids=[m.id for m in rows];db.get(Contact,rows[0].contact_id).opted_out=True;db.commit()
    process_one(ids[0]);assert calls==[]
    monkeypatch.setattr(settings,'meta_phone_number_id','999')
    process_one(ids[1]);assert calls==[]

def test_pause_racing_with_claim_preserves_pending_item(client,monkeypatch):
    cid,_,_=prepare(client);calls=fake_live(monkeypatch)
    client.post(f'/api/bulk/campaign/{cid}/start',json={'confirmed':True})
    with Session() as db:mid=db.scalar(select(Message.id))
    def pause_before_network(db):
        assert client.post(f'/api/campaigns/{cid}/pause').status_code==200
        return True
    monkeypatch.setattr('app.worker.reserve_bulk_slot',pause_before_network)
    process_one(mid)
    assert calls==[]
    with Session() as db:assert db.get(Message,mid).status=='queued';assert db.get(Message,mid).attempts==0
    client.post(f'/api/campaigns/{cid}/resume')
    monkeypatch.setattr('app.worker.reserve_bulk_slot',lambda db:True)
    process_one(mid);assert len(calls)==1

def test_remote_pending_template_prevents_enqueue(client,monkeypatch):
    cid,_,_=prepare(client);calls=fake_live(monkeypatch)
    monkeypatch.setattr('app.bulk.graph',lambda *a,**k:{'data':[{'name':'bulk_marketing','language':'tr','id':'777','status':'PENDING','category':'MARKETING'}]})
    assert client.post(f'/api/bulk/campaign/{cid}/start',json={'confirmed':True}).status_code==403
    assert calls==[]
    with Session() as db:assert db.scalar(select(Message.id)) is None

def test_prepare_replaces_stale_draft_but_reuses_unchanged_draft(client):
    cid,p,tid=prepare(client)
    data={'token':p['token'],'template_id':tid,'variables':{'1':'HYS'}}
    assert client.post('/api/bulk/campaign',json=data).json()['campaign']['id']==cid
    with Session() as db:
        row=db.get(SystemValue,f'bulk_config:{cid}');cfg=json.loads(row.value)
        cfg['definition_hash']='stale';row.value=json.dumps(cfg);db.commit()
    r=client.post('/api/bulk/campaign',json=data)
    assert r.status_code==200,r.text
    assert r.json()['campaign']['id']!=cid
    with Session() as db:assert db.scalar(select(Message.id)) is None

def test_retry_only_explicit_rejection_timeout_uncertain_and_error_export(client,monkeypatch):
    cid,_,_=prepare(client);calls=fake_live(monkeypatch)
    client.post(f'/api/bulk/campaign/{cid}/start',json={'confirmed':True})
    with Session() as db:ids=list(db.scalars(select(Message.id)))
    def reject(*a,**k):raise HTTPException(502,{'code':130429,'message':'Hız sınırı','retryable':True,'ambiguous':False})
    monkeypatch.setattr('app.worker.graph',reject);process_one(ids[0])
    with Session() as db:assert db.get(Message,ids[0]).status=='retry'
    def timeout(*a,**k):raise HTTPException(502,'Timeout')
    monkeypatch.setattr('app.worker.graph',timeout);process_one(ids[1]);process_one(ids[1])
    with Session() as db:assert db.get(Message,ids[1]).status=='uncertain';assert db.get(Message,ids[1]).attempts==1
    r=client.get(f'/api/bulk/campaign/{cid}/failures.xlsx');assert r.status_code==200
    wb=load_workbook(io.BytesIO(r.content));assert wb.active.max_row==2
    assert 'Timeout' in wb.active['C2'].value

def test_global_pacing_and_daily_cap(client,monkeypatch):
    monkeypatch.setattr('app.worker.time.time',lambda:100.)
    with Session() as db:
        assert reserve_bulk_slot(db);assert not reserve_bulk_slot(db)
        monkeypatch.setattr('app.worker.time.time',lambda:102.)
        monkeypatch.setattr(settings,'bulk_daily_limit',1)
        assert not reserve_bulk_slot(db)

def test_webhook_delivery_counts_are_exclusive(client,monkeypatch):
    cid,_,_=prepare(client);calls=fake_live(monkeypatch)
    client.post(f'/api/bulk/campaign/{cid}/start',json={'confirmed':True})
    with Session() as db:mid=db.scalar(select(Message.id))
    process_one(mid)
    body=json.dumps({'entry':[{'id':settings.meta_waba_id,'changes':[{'field':'messages','value':{'metadata':{'phone_number_id':settings.meta_phone_number_id},'statuses':[{'id':'wamid.bulk.1','status':'read'}]}}]}]}).encode()
    signature='sha256='+hmac.new(settings.meta_app_secret.encode(),body,hashlib.sha256).hexdigest()
    assert client.post('/api/webhook',content=body,headers={'X-Hub-Signature-256':signature}).status_code==200
    p=client.get(f'/api/bulk/campaign/{cid}').json()
    assert p['read']==1;assert p['accepted']==p['sent']==p['delivered']==0
