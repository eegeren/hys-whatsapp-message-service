import io,json,hmac,hashlib
from datetime import timedelta
import pytest
from sqlalchemy import select
from app.core import *
from app.worker import dispatch_local,process_one

def contact(client,kind='customer',number='0532 123 45 67'):
    r=client.post('/api/contacts/'+kind,json={'first_name':'Test','last_name':'Kayıt','phone':number,'store':'Bolu'})
    assert r.status_code==200,r.text
    return r.json()['id']
def consent(client,id,kind='customer'):
    r=client.put(f'/api/contact/{id}/consent',json={'verified':True,'scope':'marketing' if kind=='customer' else 'staff','source':'İmzalı form','evidence':'test-evidence-001','consent_date':'2026-10-08','iys_manual_verified':True,'legal_basis':'Personel iletişim bilgilendirmesi'})
    assert r.status_code==200,r.text
def campaign(client,kind='customer'):
    t=client.post('/api/templates',json={'name':'test_template','body':'Merhaba {{personel_adi}}','category':'MARKETING'}).json()
    r=client.post('/api/campaigns',json={'name':'Test','kind':kind,'template_id':t['id']})
    assert r.status_code==200,r.text
    return r.json()['id']
@pytest.mark.parametrize('value',['0532 123 45 67','5321234567','905321234567','+905321234567','00905321234567'])
def test_phone_normalization(value):assert phone(value)=='+905321234567'
@pytest.mark.parametrize('value',['abc','0000000000','123','+0001234567'])
def test_phone_reject(value):
    with pytest.raises(ValueError):phone(value)
def test_duplicates_and_module_separation(client):
    contact(client)
    assert client.post('/api/contacts/customer',json={'first_name':'X','phone':'+905321234567'}).status_code==409
    contact(client,'staff')
    assert client.get('/api/dashboard').json()['customers']==1
    assert client.get('/api/dashboard').json()['staff']==1
def test_csv_preview_never_trusts_permission(client):
    content='ad,soyad,telefon,izin\nAyşe,Eren,05321234567,izinli\nAyşe,Eren,+905321234567,izinli\nHata,Satır,abc,izinli\n'
    r=client.post('/api/import/customer/preview',files={'file':('a.csv',content.encode(),'text/csv')});assert r.status_code==200,r.text
    preview=r.json();assert preview['valid_count']==1;assert len(preview['errors'])==2
    assert client.get('/api/contacts/customer').json()['total']==0
    r=client.post('/api/import/customer/commit',json={'token':preview['token']});assert r.json()['imported']==1
    c=client.get('/api/contacts/customer').json()['items'][0];assert c['phone']=='+905321234567';assert not c['eligible'];assert c['consent']=='unverified'
    assert client.post('/api/import/customer/commit',json={'token':preview['token']}).status_code==404
def test_xlsx_staff_preview(client):
    from openpyxl import Workbook
    wb=Workbook();ws=wb.active;ws.append(['ad','telefon','mağaza','departman']);ws.append(['Test','05321234567','Bolu','Satış']);buf=io.BytesIO();wb.save(buf)
    r=client.post('/api/import/staff/preview',files={'file':('staff.xlsx',buf.getvalue(),'application/octet-stream')});assert r.status_code==200,r.text
    assert client.post('/api/import/staff/commit',json={'token':r.json()['token']}).json()['imported']==1
    assert client.get('/api/contacts/staff').json()['total']==1
def test_permission_requires_evidence_and_manual_iys(client):
    id=contact(client)
    assert client.put(f'/api/contact/{id}/consent',json={'verified':True}).status_code==422
    r=client.put(f'/api/contact/{id}/consent',json={'verified':True,'scope':'marketing','source':'form','evidence':'evidence','consent_date':'2026-10-08'});assert r.status_code==200;assert not r.json()['eligible']
    consent(client,id);assert client.get('/api/dashboard').json()['eligible']==1
def test_dry_run_and_persistence_and_idempotence(client):
    id=contact(client);c=campaign(client)
    assert client.post(f'/api/campaigns/{c}/start').status_code==422
    consent(client,id)
    assert client.get(f'/api/campaigns/{c}/preview').json()['count']==1
    assert client.post(f'/api/campaigns/{c}/start').status_code==200
    assert client.post(f'/api/campaigns/{c}/start').status_code==409
    dispatch_local();dispatch_local()
    rows=client.get('/api/messages').json()['items'];assert len(rows)==1;assert rows[0]['status']=='dry_run';assert rows[0]['meta_id'] is None
    with Session() as db:assert db.scalar(select(Message)).status=='dry_run'
def test_permission_revoked_before_send_blocks(client):
    id=contact(client);consent(client,id);c=campaign(client);client.post(f'/api/campaigns/{c}/start')
    client.put(f'/api/contact/{id}/consent',json={'verified':False,'opted_out':True})
    dispatch_local();assert client.get('/api/messages').json()['items'][0]['status']=='blocked'
def test_schedule_pause_cancel(client):
    id=contact(client);consent(client,id);c=campaign(client)
    with Session() as db:db.get(Campaign,c).scheduled=now()+timedelta(days=1);db.commit()
    client.post(f'/api/campaigns/{c}/start');dispatch_local();assert client.get('/api/messages').json()['items'][0]['status']=='queued'
    assert client.post(f'/api/campaigns/{c}/pause').status_code==200
    assert client.post(f'/api/campaigns/{c}/resume').status_code==200
    client.post(f'/api/campaigns/{c}/cancel');dispatch_local();assert client.get('/api/messages').json()['items'][0]['status']=='cancelled'
def signed(client,payload):
    raw=json.dumps(payload).encode();signature='sha256='+hmac.new(settings.meta_app_secret.encode(),raw,hashlib.sha256).hexdigest()
    return client.post('/api/webhook',content=raw,headers={'X-Hub-Signature-256':signature})
def event(messages=[],statuses=[]):return {'entry':[{'changes':[{'value':{'metadata':{'phone_number_id':'123456'},'messages':messages,'statuses':statuses}}]}]}
def test_webhook_signature_dedup_and_optout(client):
    id=contact(client);consent(client,id)
    assert client.post('/api/webhook',json={}).status_code==403
    payload=event([{'id':'wamid.in.1','from':'905321234567','type':'text','text':{'body':'STOP'}}])
    assert signed(client,payload).status_code==200;assert signed(client,payload).json()['duplicate']
    altered=event([{'id':'wamid.in.1','from':'905321234567','type':'text','text':{'body':'STOP'}}],[]);altered['other']='different body'
    signed(client,altered)
    assert client.get('/api/messages?direction=in').json()['total']==1
    c=client.get('/api/contacts/customer').json()['items'][0];assert c['opted_out'];assert not c['eligible']
def test_delivery_status_out_of_order(client):
    with Session() as db:db.add(Message(phone='+905321234567',body='Test',status='accepted',meta_id='wamid.out'));db.commit()
    for state in ['read','delivered','sent']:
        assert signed(client,event(statuses=[{'id':'wamid.out','status':state}])).status_code==200
    assert client.get('/api/messages').json()['items'][0]['status']=='read'
def test_reply_window(client):
    with Session() as db:
        m=Message(phone='+905321234567',body='Test',direction='in',status='received',created=now()-timedelta(hours=25));db.add(m);db.commit();id=m.id
    assert client.post(f'/api/messages/{id}/reply',json={'body':'Yanıt'}).status_code==403
    with Session() as db:db.get(Message,id).created=now();db.commit()
    assert client.post(f'/api/messages/{id}/reply',json={'body':'Yanıt'}).json()['status']=='dry_run'
def test_two_different_admins_and_live_lock(client):
    id=contact(client);consent(client,id);c=campaign(client)
    assert client.post(f'/api/campaigns/{c}/approve').status_code==200
    assert client.post(f'/api/campaigns/{c}/confirm').status_code==403
    settings.dry_run=False
    try:assert client.post(f'/api/campaigns/{c}/start').status_code==403
    finally:settings.dry_run=True
def test_operator_permissions_and_masking(client):
    contact(client)
    assert client.post('/api/users',json={'username':'operator','password':'TestPassword123!','role':'operator'}).status_code==200
    client.post('/api/logout');client.post('/api/login',json={'username':'operator','password':'TestPassword123!'})
    assert '•' in client.get('/api/contacts/customer').json()['items'][0]['phone']
    assert client.get('/api/settings').status_code==403
    assert client.get('/api/export/customer').status_code==403
    assert client.put('/api/contact/1/consent',json={'verified':False}).status_code==403
def test_delete_anonymizes_messages(client):
    id=contact(client);consent(client,id);c=campaign(client);client.post(f'/api/campaigns/{c}/start')
    assert client.delete(f'/api/contact/{id}').status_code==200
    with Session() as db:
        m=db.scalar(select(Message));assert m.phone=='deleted';assert m.contact_id is None;assert m.status=='cancelled'
def test_export_formula_escape(client):
    from openpyxl import load_workbook
    client.post('/api/contacts/customer',json={'first_name':'=1+1','phone':'05321234567'})
    r=client.get('/api/export/customer');assert r.status_code==200
    wb=load_workbook(io.BytesIO(r.content));assert wb.active['A2'].value=="'=1+1"
def test_csrf_and_auth(client):
    client.headers.pop('X-HYS-Request');assert client.post('/api/stores',json={'name':'Test'}).status_code==403
    client.cookies.clear();assert client.get('/api/dashboard').status_code==401
